#!/usr/bin/env python3
"""Generate deterministic binary fixtures required by DocReader tests."""

from pathlib import Path
import shutil
import subprocess

from docx import Document
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches


ROOT = Path("/workspace")
TESTDATA = ROOT / "testdata" / "rag_test"


def add_cell_paragraphs(cell, values):
    cell.text = values[0]
    for value in values[1:]:
        cell.add_paragraph(value)


def create_docx() -> Path:
    target = TESTDATA / "docx" / "en_tables.docx"
    target.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("Table normalization fixture", level=1)

    table = doc.add_table(rows=2, cols=4)
    for cell, value in zip(table.rows[0].cells, ["Name", "Game", "Fame", "Blame"]):
        cell.text = value
    for cell, value in zip(
        table.rows[1].cells, ["Lebron James", "Basketball", "Famous", "None"]
    ):
        cell.text = value

    table = doc.add_table(rows=2, cols=2)
    for cell, value in zip(table.rows[0].cells, ["Sinple", "Table"]):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ["Without", "Header"]):
        cell.text = value

    table = doc.add_table(rows=2, cols=2)
    add_cell_paragraphs(table.rows[0].cells[0], ["Simple", "Multiparagraph"])
    add_cell_paragraphs(table.rows[0].cells[1], ["Table", "Full"])
    add_cell_paragraphs(table.rows[1].cells[0], ["Of", "Paragraphs"])
    add_cell_paragraphs(table.rows[1].cells[1], ["In each", "Cell."])
    doc.save(target)
    return target


def create_presentations() -> tuple[Path, Path]:
    pptx_target = TESTDATA / "pptx" / "en_marker.pptx"
    ppt_target = TESTDATA / "ppt_old" / "en_38256.ppt"
    pptx_target.parent.mkdir(parents=True, exist_ok=True)
    ppt_target.parent.mkdir(parents=True, exist_ok=True)

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "DocReader regression fixture"
    slide.placeholders[1].text = "Modern PPTX input and legacy PPT conversion"
    image_path = TESTDATA / "pptx" / "fixture-image.png"
    image = Image.new("RGB", (320, 180), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((12, 12, 308, 168), outline="black", width=4)
    draw.ellipse((110, 40, 210, 140), fill="#10a37f")
    image.save(image_path)
    slide.shapes.add_picture(str(image_path), Inches(7.2), Inches(4.5), width=Inches(2))
    deck.save(pptx_target)
    image_path.unlink()

    profile = TESTDATA / ".libreoffice-profile"
    if profile.exists():
        shutil.rmtree(profile)
    profile.mkdir(parents=True)
    command = [
        shutil.which("soffice") or "soffice",
        "--headless",
        f"-env:UserInstallation={profile.as_uri()}",
        "--convert-to",
        "ppt:MS PowerPoint 97",
        "--outdir",
        str(ppt_target.parent),
        str(pptx_target),
    ]
    result = subprocess.run(command, text=True, capture_output=True, timeout=120)
    converted = ppt_target.parent / f"{pptx_target.stem}.ppt"
    if result.returncode == 0 and converted.is_file() and converted != ppt_target:
        converted.replace(ppt_target)
    if result.returncode != 0 or not ppt_target.is_file():
        raise RuntimeError(
            f"LibreOffice conversion failed: exit={result.returncode} "
            f"stdout={result.stdout!r} stderr={result.stderr!r}"
        )
    shutil.rmtree(profile)
    return pptx_target, ppt_target


if __name__ == "__main__":
    paths = (create_docx(), *create_presentations())
    for path in paths:
        print(f"FIXTURE {path.relative_to(ROOT)} bytes={path.stat().st_size}")
